import argparse
import os
import re
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation

import openpyxl
import pyodbc


# ============================================================
# CONFIGURACIÓN
# ============================================================

SERVER = r"LAPTOP-6VL1O8NQ\SQLEXPRESS"
DATABASE = "SNCEV"

CONNECTION_STRING = (
    "DRIVER={ODBC Driver 18 for SQL Server};"
    f"SERVER={SERVER};"
    f"DATABASE={DATABASE};"
    "Trusted_Connection=yes;"
    "TrustServerCertificate=yes;"
)

TABLE = "sncev_geo.establecimiento_salud"


# ============================================================
# ENCABEZADOS OFICIALES
# ============================================================

EXPECTED_HEADERS = [
    "CLUES",
    "CLAVE DE LA INSTITUCION",
    "NOMBRE DE LA INSTITUCION",
    "CLAVE DE LA ENTIDAD",
    "ENTIDAD",
    "CLAVE DEL MUNICIPIO",
    "MUNICIPIO",
    "CLAVE DE LA LOCALIDAD",
    "LOCALIDAD",
    "CLAVE DE LA JURISDICCION",
    "JURISDICCION",
    "CLAVE DEL TIPO ESTABLECIMIENTO",
    "NOMBRE TIPO ESTABLECIMIENTO",
    "CLAVE DE TIPOLOGIA",
    "NOMBRE DE TIPOLOGIA",
    "CLAVE DE SUBTIPOLOGIA",
    "NOMBRE DE SUBTIPOLOGIA",
    "NOMBRE DE LA UNIDAD",
    "NOMBRE COMERCIAL",
    "CLAVE TIPO DE VIALIDAD",
    "TIPO DE VIALIDAD",
    "VIALIDAD",
    "NUMERO EXTERIOR",
    "NUMERO INTERIOR",
    "CLAVE TIPO DE ASENTAMIENTO",
    "TIPO DE ASENTAMIENTO",
    "ASENTAMIENTO",
    "CODIGO POSTAL",
    "REFERENCIAS DEL DOMICILIO",
    "CLAVE ESTATUS DE OPERACION",
    "ESTATUS DE OPERACION",
    "RFC DEL ESTABLECIMIENTO",
    "TELEFONO 1 DEL ESTABLECIMIENTO",
    "EXTENSION TELEFONICA 1 DEL ESTABLECIMIENTO",
    "TELEFONO 2 DEL ESTABLECIMIENTO",
    "EXTENSION TELEFONICA 2 DEL ESTABLECIMIENTO",
    "FECHA DE CONSTRUCCION",
    "FECHA DE INICIO DE OPERACION",
    "CLAVE UNIDAD MOVIL MARCA",
    "UNIDAD MOVIL MARCA",
    "UNIDAD MOVIL MARCA ESPECIFICA",
    "UNIDAD MOVIL MODELO",
    "CLAVE UNIDAD MOVIL PROGRAMA",
    "UNIDAD MOVIL PROGRAMA",
    "CLAVE UNIDAD MOVIL TIPO",
    "UNIDAD MOVIL TIPO",
    "CLAVE UNIDAD MOVIL TIPOLOGIA",
    "UNIDAD MOVIL TIPOLOGIA",
    "CLAVE DE LA INS ADM",
    "NOMBRE DE LA INS ADM",
    "CLAVE NIVEL ATENCION",
    "NIVEL ATENCION",
    "CLAVE ESTRATO UNIDAD",
    "ESTRATO UNIDAD",
    "CLAVE TIPO OBRA",
    "TIPO OBRA",
    "CLAVE PROPIEDAD DEL INMUEBLE",
    "PROPIEDAD DEL INMUEBLE",
    "OBSERVACIONES AL REGISTRO",
    "LATITUD",
    "LONGITUD",
    "CLAVE ULTIMO MOVIMIENTO",
    "ULTIMO MOVIMIENTO",
    "FECHA ULTIMO MOVIMIENTO",
    "COMENTARIOS DE LA VALIDACION",
    "CLAVE MOTIVO BAJA",
    "MOTIVO BAJA",
    "FECHA EFECTIVA DE BAJA",
]


DB_COLUMNS = [
    "clues",
    "clave_institucion",
    "nombre_institucion",
    "clave_entidad",
    "entidad",
    "clave_municipio",
    "municipio",
    "clave_localidad",
    "localidad",
    "clave_jurisdiccion",
    "jurisdiccion",
    "clave_tipo_establecimiento",
    "nombre_tipo_establecimiento",
    "clave_tipologia",
    "nombre_tipologia",
    "clave_subtipologia",
    "nombre_subtipologia",
    "nombre_unidad",
    "nombre_comercial",
    "clave_tipo_vialidad",
    "tipo_vialidad",
    "vialidad",
    "numero_exterior",
    "numero_interior",
    "clave_tipo_asentamiento",
    "tipo_asentamiento",
    "asentamiento",
    "codigo_postal",
    "referencias_domicilio",
    "clave_estatus_operacion",
    "estatus_operacion",
    "rfc_establecimiento",
    "telefono_1",
    "extension_telefonica_1",
    "telefono_2",
    "extension_telefonica_2",
    "fecha_construccion",
    "fecha_inicio_operacion",
    "clave_unidad_movil_marca",
    "unidad_movil_marca",
    "unidad_movil_marca_especifica",
    "unidad_movil_modelo",
    "clave_unidad_movil_programa",
    "unidad_movil_programa",
    "clave_unidad_movil_tipo",
    "unidad_movil_tipo",
    "clave_unidad_movil_tipologia",
    "unidad_movil_tipologia",
    "clave_ins_adm",
    "nombre_ins_adm",
    "clave_nivel_atencion",
    "nivel_atencion",
    "clave_estrato_unidad",
    "estrato_unidad",
    "clave_tipo_obra",
    "tipo_obra",
    "clave_propiedad_inmueble",
    "propiedad_inmueble",
    "observaciones_registro",
    "latitud",
    "longitud",
    "clave_ultimo_movimiento",
    "ultimo_movimiento",
    "fecha_ultimo_movimiento",
    "comentarios_validacion",
    "clave_motivo_baja",
    "motivo_baja",
    "fecha_efectiva_baja",
]


# ============================================================
# CONVERSIÓN DE DATOS
# ============================================================

def texto(valor):
    if valor is None:
        return None

    valor = str(valor).strip()

    if valor == "":
        return None

    return valor


def fecha(valor):
    if valor is None:
        return None

    if isinstance(valor, datetime):
        return valor.date()

    s = str(valor).strip()

    if not s:
        return None

    formatos = (
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%Y/%m/%d",
    )

    for formato in formatos:
        try:
            return datetime.strptime(s, formato).date()
        except ValueError:
            pass

    raise ValueError(f"Fecha no reconocida: {s!r}")


def decimal_nullable(valor):
    if valor is None:
        return None

    s = str(valor).strip()

    if not s:
        return None

    s = s.replace(",", ".")

    try:
        return Decimal(s)
    except InvalidOperation:
        raise ValueError(f"Decimal no válido: {valor!r}")


def normalizar_fila(row):
    valores = [texto(v) for v in row]

    # Fechas
    valores[36] = fecha(row[36])
    valores[37] = fecha(row[37])
    valores[63] = fecha(row[63])
    valores[67] = fecha(row[67])

    # Coordenadas
    valores[59] = decimal_nullable(row[59])
    valores[60] = decimal_nullable(row[60])

    return valores


def es_seleccionable(valores):
    """
    No borra históricos.

    Una unidad sólo será seleccionable inicialmente cuando el
    catálogo la declare EN OPERACION y no tenga fecha efectiva
    de baja.

    Más adelante podremos endurecer esta política según
    institución/tipología/nivel/servicios de vacunación.
    """

    estatus = (valores[30] or "").strip().upper()
    fecha_baja = valores[67]

    return 1 if estatus == "EN OPERACION" and fecha_baja is None else 0


# ============================================================
# VERSIÓN DEL CATÁLOGO
# ============================================================

def obtener_version(nombre_archivo, nombre_hoja):
    candidatos = [
        os.path.basename(nombre_archivo),
        nombre_hoja,
    ]

    for candidato in candidatos:
        match = re.search(r"(20\d{4})", candidato)

        if match:
            return match.group(1)

    raise ValueError(
        "No fue posible obtener la versión YYYYMM "
        "del nombre del archivo o de la hoja."
    )


# ============================================================
# VALIDACIÓN DEL EXCEL
# ============================================================

def validar_encabezados(headers):
    if len(headers) != len(EXPECTED_HEADERS):
        raise RuntimeError(
            f"El archivo contiene {len(headers)} columnas; "
            f"se esperaban {len(EXPECTED_HEADERS)}."
        )

    errores = []

    for i, (actual, esperado) in enumerate(
        zip(headers, EXPECTED_HEADERS),
        start=1,
    ):
        actual = "" if actual is None else str(actual).strip()

        if actual != esperado:
            errores.append(
                f"Columna {i}: "
                f"esperada={esperado!r}, recibida={actual!r}"
            )

    if errores:
        detalle = "\n".join(errores[:20])

        raise RuntimeError(
            "La estructura del catálogo CLUES cambió.\n"
            + detalle
        )


# ============================================================
# SQL
# ============================================================

def crear_sql_update():
    sets = [
        f"{col}=?"
        for col in DB_COLUMNS[1:]
    ]

    sets.extend([
        "seleccionable=?",
        "version_catalogo=?",
        "fecha_importacion=SYSDATETIME()",
    ])

    return f"""
        UPDATE {TABLE}
        SET {", ".join(sets)}
        WHERE clues=?
    """


def crear_sql_insert():
    columnas = DB_COLUMNS + [
        "seleccionable",
        "version_catalogo",
    ]

    placeholders = ",".join(
        ["?"] * len(columnas)
    )

    return f"""
        INSERT INTO {TABLE}
        (
            {", ".join(columnas)}
        )
        VALUES
        (
            {placeholders}
        )
    """


# ============================================================
# IMPORTACIÓN
# ============================================================

def importar(ruta):
    ruta = os.path.abspath(ruta)

    if not os.path.isfile(ruta):
        raise FileNotFoundError(ruta)

    print("=" * 72)
    print("SNCEV - IMPORTADOR NACIONAL CLUES")
    print("=" * 72)
    print(f"Archivo : {ruta}")

    workbook = openpyxl.load_workbook(
        ruta,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.worksheets[0]

    print(f"Hoja    : {worksheet.title}")
    print(f"Filas   : {worksheet.max_row}")
    print(f"Columnas: {worksheet.max_column}")

    headers = [
        celda.value
        for celda in next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
            )
        )
    ]

    validar_encabezados(headers)

    version = obtener_version(
        ruta,
        worksheet.title,
    )

    print(f"Versión : {version}")
    print("Estructura de 68 columnas: OK")
    print()

    conn = pyodbc.connect(
        CONNECTION_STRING,
        autocommit=False,
    )

    cursor = conn.cursor()

    importacion_id = None

    leidos = 0
    insertados = 0
    actualizados = 0
    sin_cambios = 0
    errores = 0

    try:
        # --------------------------------------------------------
        # Auditoría de importación
        # --------------------------------------------------------

        cursor.execute(
            """
            INSERT INTO sncev_geo.importacion_clues
            (
                version_catalogo,
                nombre_archivo,
                estado
            )
            OUTPUT INSERTED.id
            VALUES (?, ?, 'INICIADA')
            """,
            version,
            os.path.basename(ruta),
        )

        importacion_id = cursor.fetchone()[0]

        conn.commit()

        sql_insert = crear_sql_insert()
        sql_update = crear_sql_update()

        # Cargamos las CLUES existentes una sola vez.
        cursor.execute(
            f"""
            SELECT clues
            FROM {TABLE}
            """
        )

        existentes = {
            str(row[0]).strip()
            for row in cursor.fetchall()
        }

        print(
            f"CLUES existentes en SNCEV: "
            f"{len(existentes):,}"
        )
        print()
        print("Importando...")

        for numero_fila, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            leidos += 1

            try:
                valores = normalizar_fila(row)

                clues = valores[0]

                if not clues:
                    raise ValueError(
                        "Registro sin CLUES."
                    )

                seleccionable = es_seleccionable(
                    valores
                )

                if clues in existentes:

                    # --------------------------------------------
                    # ¿Cambió realmente?
                    # --------------------------------------------

                    cursor.execute(
                        f"""
                        SELECT
                            {", ".join(DB_COLUMNS[1:])},
                            seleccionable
                        FROM {TABLE}
                        WHERE clues=?
                        """,
                        clues,
                    )

                    actual = cursor.fetchone()

                    nuevo_comparable = tuple(
                        valores[1:]
                    ) + (bool(seleccionable),)

                    actual_comparable = tuple(actual)

                    if actual_comparable == nuevo_comparable:
                        sin_cambios += 1

                        # Aunque no cambió el contenido,
                        # actualizamos versión del catálogo.
                        cursor.execute(
                            f"""
                            UPDATE {TABLE}
                            SET
                                version_catalogo=?,
                                fecha_importacion=
                                    SYSDATETIME()
                            WHERE clues=?
                            """,
                            version,
                            clues,
                        )

                    else:
                        parametros = (
                            valores[1:]
                            + [
                                seleccionable,
                                version,
                                clues,
                            ]
                        )

                        cursor.execute(
                            sql_update,
                            parametros,
                        )

                        actualizados += 1

                else:
                    parametros = (
                        valores
                        + [
                            seleccionable,
                            version,
                        ]
                    )

                    cursor.execute(
                        sql_insert,
                        parametros,
                    )

                    existentes.add(clues)
                    insertados += 1

                # Commit periódico.
                if leidos % 1000 == 0:
                    conn.commit()

                    print(
                        f"  {leidos:>7,} / "
                        f"{worksheet.max_row - 1:>7,} "
                        f"| nuevos {insertados:,} "
                        f"| actualizados {actualizados:,} "
                        f"| iguales {sin_cambios:,}"
                    )

            except Exception as ex:
                errores += 1

                print(
                    f"\nERROR fila {numero_fila}: "
                    f"{ex}",
                    file=sys.stderr,
                )

                # Política fail-fast.
                raise

        conn.commit()

        # --------------------------------------------------------
        # Finalizar auditoría
        # --------------------------------------------------------

        cursor.execute(
            """
            UPDATE sncev_geo.importacion_clues
            SET
                fecha_fin=SYSDATETIME(),
                registros_leidos=?,
                registros_insertados=?,
                registros_actualizados=?,
                registros_sin_cambios=?,
                registros_error=?,
                estado='COMPLETADA',
                mensaje='Importación finalizada correctamente.'
            WHERE id=?
            """,
            leidos,
            insertados,
            actualizados,
            sin_cambios,
            errores,
            importacion_id,
        )

        conn.commit()

        print()
        print("=" * 72)
        print("IMPORTACIÓN COMPLETADA")
        print("=" * 72)
        print(f"Leídos       : {leidos:,}")
        print(f"Insertados   : {insertados:,}")
        print(f"Actualizados : {actualizados:,}")
        print(f"Sin cambios  : {sin_cambios:,}")
        print(f"Errores      : {errores:,}")
        print(f"Versión      : {version}")
        print("=" * 72)

    except Exception as ex:
        conn.rollback()

        if importacion_id is not None:
            try:
                cursor.execute(
                    """
                    UPDATE sncev_geo.importacion_clues
                    SET
                        fecha_fin=SYSDATETIME(),
                        registros_leidos=?,
                        registros_insertados=?,
                        registros_actualizados=?,
                        registros_sin_cambios=?,
                        registros_error=?,
                        estado='ERROR',
                        mensaje=?
                    WHERE id=?
                    """,
                    leidos,
                    insertados,
                    actualizados,
                    sin_cambios,
                    errores,
                    str(ex)[:2000],
                    importacion_id,
                )

                conn.commit()

            except Exception:
                conn.rollback()

        raise

    finally:
        workbook.close()
        cursor.close()
        conn.close()


# ============================================================
# MAIN
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description=(
            "Importa el catálogo nacional de "
            "establecimientos CLUES a SNCEV."
        )
    )

    parser.add_argument(
        "archivo",
        help="Ruta del archivo ESTABLECIMIENTO_SALUD_YYYYMM.xlsx",
    )

    args = parser.parse_args()

    importar(args.archivo)


if __name__ == "__main__":
    main()